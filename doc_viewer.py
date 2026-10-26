"""Built-in document viewer — opens PDFs, DOCX, and TXT cross-platform."""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

from PyQt6.QtCore import Qt, QRectF
from PyQt6.QtGui import QImage, QPixmap, QPainter, QColor, QPen
from PyQt6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QScrollArea,
    QLabel, QPushButton, QWidget, QPlainTextEdit,
)


class DocViewerDialog(QDialog):
    """Modal viewer that opens a document at a specific page and highlights a word."""

    def __init__(self, filepath: str, pagenumber: int, word: str, sentence: str = "", parent=None):
        super().__init__(parent)
        self.setWindowTitle(Path(filepath).name)
        self.resize(860, 680)

        ext = Path(filepath).suffix.lower()
        layout = QVBoxLayout(self)

        if ext == ".pdf":
            widget = _PdfViewer(filepath, pagenumber, word)
        elif ext == ".docx":
            widget = _TextViewer(filepath, word, mode="docx", sentence=sentence)
        else:
            widget = _TextViewer(filepath, word, mode="txt", sentence=sentence)

        layout.addWidget(widget)

        close = QPushButton("Close")
        close.clicked.connect(self.accept)
        layout.addWidget(close, alignment=Qt.AlignmentFlag.AlignRight)


class _PdfViewer(QWidget):
    """Renders the target page as an image with the search word highlighted."""

    def __init__(self, filepath: str, pagenumber: int, word: str):
        super().__init__()
        layout = QVBoxLayout(self)

        # Navigation bar
        nav = QHBoxLayout()
        self._page_label = QLabel()
        self._prev = QPushButton("◀ Prev")
        self._next = QPushButton("Next ▶")
        self._prev.clicked.connect(self._go_prev)
        self._next.clicked.connect(self._go_next)
        nav.addWidget(self._prev)
        nav.addWidget(self._page_label, 1, Qt.AlignmentFlag.AlignCenter)
        nav.addWidget(self._next)
        layout.addLayout(nav)

        self._scroll = QScrollArea()
        self._scroll.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._img_label = QLabel()
        self._img_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._scroll.setWidget(self._img_label)
        self._scroll.setWidgetResizable(True)
        layout.addWidget(self._scroll)

        try:
            import fitz
            self._doc = fitz.open(filepath)
        except Exception as e:
            self._img_label.setText(f"Cannot open PDF: {e}")
            return

        self._word = word
        self._total = len(self._doc)
        self._current = max(0, min(pagenumber - 1, self._total - 1))
        self._render()

    def _render(self):
        import fitz
        page = self._doc[self._current]
        mat = fitz.Matrix(1.8, 1.8)
        pix = page.get_pixmap(matrix=mat)

        img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888)
        pixmap = QPixmap.fromImage(img)

        # Highlight all occurrences of the word on this page
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        pen = QPen(QColor(255, 220, 0, 200))
        pen.setWidth(2)
        painter.setPen(pen)
        painter.setBrush(QColor(255, 220, 0, 60))
        scale = 1.8
        rects = [r for w in (self._word.split() or [self._word]) for r in page.search_for(w)]
        for rect in rects:
            r = QRectF(
                rect.x0 * scale, rect.y0 * scale,
                (rect.x1 - rect.x0) * scale, (rect.y1 - rect.y0) * scale,
            )
            painter.drawRect(r)
        painter.end()

        self._img_label.setPixmap(pixmap)
        self._img_label.resize(pixmap.size())
        self._page_label.setText(f"Page {self._current + 1} / {self._total}")
        self._prev.setEnabled(self._current > 0)
        self._next.setEnabled(self._current < self._total - 1)

    def _go_prev(self):
        self._current -= 1
        self._render()

    def _go_next(self):
        self._current += 1
        self._render()


class _TextViewer(QWidget):
    """Shows plain text content with the search word highlighted."""

    def __init__(self, filepath: str, word: str, mode: str, sentence: str = ""):
        super().__init__()
        layout = QVBoxLayout(self)

        editor = QPlainTextEdit()
        editor.setReadOnly(True)
        editor.setStyleSheet("font-family: monospace; font-size: 11px;")
        layout.addWidget(editor)

        try:
            if mode == "docx":
                from docx import Document
                doc = Document(filepath)
                text = "\n".join(p.text for p in doc.paragraphs)
            else:
                text = Path(filepath).read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            editor.setPlainText(f"Cannot open file: {e}")
            return

        editor.setPlainText(text)

        # Scroll to the specific sentence if given, otherwise to the first word occurrence
        target = sentence[:80] if sentence else word
        cursor = editor.document().find(target)
        if cursor.isNull() and sentence:
            cursor = editor.document().find(word.split()[0] if word.split() else word)
        if not cursor.isNull():
            editor.setTextCursor(cursor)
            editor.ensureCursorVisible()


# ---------------------------------------------------------------------------
# Sentence extraction (used by the sentence panel in app.py)
# ---------------------------------------------------------------------------

def sentences_containing(filepath: str, word: str) -> list[tuple[str, int]]:
    """Return (sentence, page_number) pairs for all sentences containing word.

    page_number is 1-based for PDFs; 0 for DOCX/TXT (page concept not applicable).
    """
    ext = Path(filepath).suffix.lower()
    w = word.lower()
    results: list[tuple[str, int]] = []

    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(filepath)
            for page_idx, page in enumerate(doc):
                text = re.sub(r'\s+', ' ', page.get_text("text"))
                for s in re.split(r'(?<=[.!?])\s+', text):
                    if s.strip() and w in s.lower():
                        results.append((s.strip(), page_idx + 1))
            doc.close()
        else:
            if ext == ".docx":
                from docx import Document
                doc = Document(filepath)
                text = "\n".join(p.text for p in doc.paragraphs)
            else:
                text = Path(filepath).read_text(encoding="utf-8", errors="replace")
            text = re.sub(r'\s+', ' ', text)
            for s in re.split(r'(?<=[.!?])\s+', text):
                if s.strip() and w in s.lower():
                    results.append((s.strip(), 0))
    except Exception:
        return []

    return results


def contexts_containing(filepath: str, words: list[str]) -> list[tuple[str, int, str]]:
    """Return (text, page_number, kind) for two search words.

    kind 'sentence': a sentence that contains both words.
    kind 'paragraph': a paragraph in which the words occur in different
    sentences; text is those sentences joined with ' … '.

    Paragraphs are split at blank lines, as allmydox does when it fills
    noun_paragraph.  page_number is 1-based for PDFs, 0 otherwise.
    """
    w1, w2 = (w.lower() for w in words[:2])
    ext = Path(filepath).suffix.lower()
    results: list[tuple[str, int, str]] = []

    def scan(text: str, page: int):
        for para in re.split(r'\n\s*\n', text):
            para = re.sub(r'\s+', ' ', para).strip()
            low = para.lower()
            if not para or w1 not in low or w2 not in low:
                continue
            sents = [x.strip() for x in re.split(r'(?<=[.!?])\s+', para) if x.strip()]
            has1 = [w1 in x.lower() for x in sents]
            has2 = [w2 in x.lower() for x in sents]
            for x, a, b in zip(sents, has1, has2):
                if a and b:
                    results.append((x, page, "sentence"))
            # Words in two different sentences of this paragraph
            if any(a and b2 for i, a in enumerate(has1)
                   for j, b2 in enumerate(has2) if i != j):
                excerpt = " … ".join(x for x, a, b in zip(sents, has1, has2) if a or b)
                results.append((excerpt, page, "paragraph"))

    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(filepath)
            for page_idx, page in enumerate(doc):
                scan(page.get_text("text"), page_idx + 1)
            doc.close()
        elif ext == ".docx":
            from docx import Document
            scan("\n".join(p.text for p in Document(filepath).paragraphs), 0)
        else:
            scan(Path(filepath).read_text(encoding="utf-8", errors="replace"), 0)
    except Exception:
        return []
    return results
